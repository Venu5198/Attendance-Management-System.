import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_sample_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees Roster"

    # Define headers
    headers = [
        "employee_id",
        "full_name",
        "email",
        "department",
        "role",
        "hire_date"
    ]
    ws.append(headers)

    # Sample realistic test employees
    employees = [
        ("EMP101", "Alexander Wright", "alexander.wright@company.com", "Engineering", "EMPLOYEE", "2024-03-15"),
        ("EMP102", "Sophia Martinez", "sophia.martinez@company.com", "Product", "EMPLOYEE", "2024-05-01"),
        ("EMP103", "Liam O'Connor", "liam.oconnor@company.com", "DevOps", "EMPLOYEE", "2024-06-10"),
        ("EMP104", "Emma Watson", "emma.watson@company.com", "Marketing", "EMPLOYEE", "2024-07-22"),
        ("EMP105", "Noah Patel", "noah.patel@company.com", "Finance", "EMPLOYEE", "2024-08-01"),
        ("EMP106", "Olivia Kim", "olivia.kim@company.com", "Human Resources", "EMPLOYEE", "2024-09-12"),
        ("EMP107", "James Wilson", "james.wilson@company.com", "Engineering", "EMPLOYEE", "2024-10-05"),
        ("EMP108", "Ava Chen", "ava.chen@company.com", "Quality Assurance", "EMPLOYEE", "2024-11-18"),
        ("EMP109", "Lucas Silva", "lucas.silva@company.com", "Design", "EMPLOYEE", "2025-01-08"),
        ("EMP110", "Mia Robinson", "mia.robinson@company.com", "Operations", "EMPLOYEE", "2025-02-14"),
    ]

    for emp in employees:
        ws.append(list(emp))

    # Styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # Apply header styles
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Apply data styles
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=len(employees) + 1, min_col=1, max_col=len(headers)), start=2):
        row_fill = PatternFill(start_color="F8FAFC" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
        for col_idx, cell in enumerate(row, start=1):
            cell.font = data_font
            cell.fill = row_fill
            cell.border = thin_border
            if col_idx in [1, 5, 6]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    from pathlib import Path

    repo_root = Path(__file__).resolve().parent.parent
    output_path = repo_root / "employees_bulk_import_sample.xlsx"
    wb.save(str(output_path))
    print(f"[*] Successfully saved XLSX to {output_path}")

if __name__ == "__main__":
    generate_sample_xlsx()
