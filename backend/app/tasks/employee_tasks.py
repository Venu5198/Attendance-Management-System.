"""Celery background tasks for employee management and bulk ingestion."""

import base64
import csv
import io
import logging
from datetime import datetime, timezone
from typing import Any

import openpyxl
import pymongo
from pymongo.uri_parser import parse_uri

from app.celery_app import celery_app
from app.config import get_settings
from app.utils.security import get_password_hash

logger = logging.getLogger(__name__)


def _get_sync_db():
    """Create synchronous MongoDB connection for Celery task worker."""
    settings = get_settings()
    from app.database import db_manager
    if db_manager.db is not None:
        db_name = db_manager.db.name
    else:
        db_name = settings.MONGODB_DATABASE
        try:
            parsed = parse_uri(settings.MONGODB_URI)
            if parsed.get("database"):
                db_name = parsed["database"]
        except Exception:
            pass
    client = pymongo.MongoClient(settings.MONGODB_URI, serverSelectionTimeoutMS=5000)
    return client[db_name]


@celery_app.task(bind=True, name="app.tasks.employee_tasks.process_employee_excel_import")
def process_employee_excel_import(
    self,
    file_content_b64: str,
    filename: str,
    tenant_id: str = "default",
    uploaded_by: str = "system",
) -> dict[str, Any]:
    """
    Process employee records from uploaded Excel (.xlsx) or CSV file in the background.

    - Updates task progress dynamically via RabbitMQ -> Celery state.
    - Saves parsed employee profiles into MongoDB with strict tenant_id scoping.
    - Generates corresponding user accounts with hashed credentials.
    - Comprehensive error handling and detailed audit report.
    """
    logger.info(
        "Celery worker received bulk import task [%s] for file '%s' (tenant: %s, user: %s)",
        self.request.id,
        filename,
        tenant_id,
        uploaded_by,
    )

    db = _get_sync_db()
    file_bytes = base64.b64decode(file_content_b64)
    rows_to_process: list[dict[str, Any]] = []

    # 1. Parse Excel or CSV
    try:
        lower_name = filename.lower()
        if lower_name.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb.active
            iter_rows = list(sheet.iter_rows(values_only=True))
            if not iter_rows:
                return {
                    "status": "FAILED",
                    "error": "The uploaded Excel file contains no data.",
                    "total_rows": 0,
                }
            raw_headers = [str(h).strip().lower() if h is not None else "" for h in iter_rows[0]]
            for row_idx, row_values in enumerate(iter_rows[1:], start=2):
                if not any(row_values):
                    continue
                row_dict = {}
                for h, val in zip(raw_headers, row_values):
                    if h:
                        row_dict[h] = str(val).strip() if val is not None else ""
                row_dict["_row_number"] = row_idx
                rows_to_process.append(row_dict)

        elif lower_name.endswith(".csv"):
            text_stream = io.StringIO(file_bytes.decode("utf-8", errors="replace"))
            reader = csv.DictReader(text_stream)
            for idx, r in enumerate(reader, start=2):
                clean_r = {k.strip().lower(): v.strip() for k, v in r.items() if k}
                clean_r["_row_number"] = idx
                rows_to_process.append(clean_r)
        else:
            return {
                "status": "FAILED",
                "error": f"Unsupported file extension for '{filename}'. Expected .xlsx or .csv",
                "total_rows": 0,
            }
    except Exception as exc:
        logger.error("Failed to parse file %s: %s", filename, exc, exc_info=True)
        return {
            "status": "FAILED",
            "error": f"Failed to parse file: {str(exc)}",
            "total_rows": 0,
        }

    total_rows = len(rows_to_process)
    if total_rows == 0:
        return {
            "status": "COMPLETED",
            "message": "No valid data rows found in uploaded file.",
            "total_rows": 0,
            "inserted_count": 0,
            "updated_count": 0,
            "failed_count": 0,
            "failed_rows": [],
        }

    # Initial state update
    if getattr(self, "request", None) and self.request.id:
        self.update_state(
            state="PROGRESS",
            meta={
                "current": 0,
                "total": total_rows,
                "percent": 0,
                "status": "Parsing complete. Ingesting records into database...",
            },
        )

    inserted_count = 0
    updated_count = 0
    failed_rows: list[dict[str, Any]] = []

    # 2. Ingest records with progress reporting
    for idx, row in enumerate(rows_to_process, start=1):
        row_num = row.get("_row_number", idx)
        emp_id = row.get("employee_id") or row.get("empid") or row.get("id")
        full_name = row.get("full_name") or row.get("name")
        email = row.get("email")
        department = row.get("department") or row.get("dept") or "Engineering"
        role = row.get("role") or "EMPLOYEE"
        hire_date = row.get("hire_date")

        # Basic validations
        if not emp_id or not full_name or not email:
            failed_rows.append({
                "row": row_num,
                "employee_id": emp_id or "MISSING",
                "error": "Missing required fields (employee_id, full_name, or email)",
            })
            continue

        emp_id = str(emp_id).strip()
        email = str(email).strip().lower()
        full_name = str(full_name).strip()
        role = role.upper() if role.upper() in ["ADMIN", "EMPLOYEE"] else "EMPLOYEE"

        try:
            now_iso = datetime.now(timezone.utc).isoformat()
            emp_doc = {
                "tenant_id": tenant_id,
                "employee_id": emp_id,
                "full_name": full_name,
                "email": email,
                "department": department,
                "role": role,
                "hire_date": hire_date or None,
                "is_active": True,
                "updated_at": now_iso,
            }

            # Upsert into employees collection
            emp_result = db["employees"].update_one(
                {"tenant_id": tenant_id, "employee_id": emp_id},
                {
                    "$set": emp_doc,
                    "$setOnInsert": {"created_at": now_iso},
                },
                upsert=True,
            )

            if emp_result.upserted_id:
                inserted_count += 1
            else:
                updated_count += 1

            # Ensure corresponding user account exists
            user_exists = db["users"].find_one({
                "$or": [
                    {"tenant_id": tenant_id, "username": emp_id},
                    {"tenant_id": tenant_id, "email": email},
                ]
            })
            if not user_exists:
                default_hashed_pwd = get_password_hash("EmpPass123!")
                db["users"].insert_one({
                    "tenant_id": tenant_id,
                    "username": emp_id,
                    "email": email,
                    "hashed_password": default_hashed_pwd,
                    "role": role,
                    "employee_id": emp_id,
                    "is_active": True,
                    "created_at": now_iso,
                    "updated_at": now_iso,
                })

        except Exception as exc:
            logger.warning("Failed to insert row %d (%s): %s", row_num, emp_id, exc)
            failed_rows.append({
                "row": row_num,
                "employee_id": emp_id,
                "error": str(exc),
            })

        # Update progress every 5 rows or on final row
        if idx % 5 == 0 or idx == total_rows:
            percent = int((idx / total_rows) * 100)
            if getattr(self, "request", None) and self.request.id:
                self.update_state(
                    state="PROGRESS",
                    meta={
                        "current": idx,
                        "total": total_rows,
                        "percent": percent,
                        "status": f"Processed {idx} of {total_rows} records ({inserted_count} created, {updated_count} updated)",
                    },
                )

    result_summary = {
        "status": "COMPLETED",
        "tenant_id": tenant_id,
        "filename": filename,
        "uploaded_by": uploaded_by,
        "total_rows": total_rows,
        "inserted_count": inserted_count,
        "updated_count": updated_count,
        "failed_count": len(failed_rows),
        "failed_rows": failed_rows,
        "completed_at": datetime.now(timezone.utc).isoformat(),
    }
    logger.info("Celery bulk import task [%s] finished: %s", self.request.id, result_summary)
    return result_summary
